"""T2 (Member A): .docx extraction including tables and text boxes."""
from docx import Document

from src.extract_docx import extract_paragraphs, extract_paragraphs_from_xml

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


def test_body_and_table_text(tmp_path):
    doc = Document()
    doc.add_paragraph("Isigidimi sama-Xosa")
    doc.add_paragraph("")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "abantu bonke"
    table.cell(0, 1).text = "umntu uthetha"
    path = tmp_path / "e.docx"
    doc.save(path)
    rows = extract_paragraphs(path)
    texts = [r["text"] for r in rows]
    assert "Isigidimi sama-Xosa" in texts and "" in texts
    in_table = {r["text"]: r["in_table"] for r in rows}
    assert in_table["abantu bonke"] and in_table["umntu uthetha"]
    assert not in_table["Isigidimi sama-Xosa"]


def test_textbox_not_duplicated_and_tabs_breaks():
    xml = f"""<w:document xmlns:w="{W}" xmlns:mc="{MC}" xmlns:wps="{WPS}"><w:body>
      <w:p><w:r><w:t>Outer</w:t></w:r>
        <mc:AlternateContent>
          <mc:Choice Requires="wps"><w:r><w:drawing><wps:txbx><w:txbxContent>
            <w:p><w:r><w:t>Ilizwe</w:t><w:tab/><w:t>lethu</w:t><w:br/><w:t>ngoku</w:t></w:r></w:p>
          </w:txbxContent></wps:txbx></w:drawing></w:r></mc:Choice>
          <mc:Fallback><w:r><w:pict><w:txbxContent>
            <w:p><w:r><w:t>Ilizwe lethu ngoku</w:t></w:r></w:p>
          </w:txbxContent></w:pict></w:r></mc:Fallback>
        </mc:AlternateContent>
      </w:p></w:body></w:document>""".encode()
    rows = extract_paragraphs_from_xml(xml)
    texts = [r["text"] for r in rows]
    assert texts.count("Ilizwe lethu\nngoku") == 1           # once, not twice
    assert "Ilizwe lethu ngoku" not in texts                  # fallback removed
    assert texts[0] == "Outer"                                # nested text not in outer paragraph
    box = [r for r in rows if r["text"].startswith("Ilizwe")][0]
    assert box["in_textbox"] and not box["in_table"]
