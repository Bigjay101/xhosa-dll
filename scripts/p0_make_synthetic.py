"""Synthetic smoke data for development BEFORE the real data is ready (Member A).

Writes fake .docx editions and fake modern caches into data/smoke/ with the same
layout as the real smoke kit, so B and C can build and test their code from day
one, and `bash run_all.sh --smoke` runs end to end.

THE TEXT IS SYNTHETIC WORD SALAD built from common isiXhosa words; historical
editions use older spellings (e.g. 'teta' for 'thetha'). It is for testing code
only and must never be used for results.
"""
import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402

from src.config import P, base_parser, load_config  # noqa: E402

MODERN = ("abantu umntu ukuba kwaye kodwa ngoku kakhulu phantsi phezulu thetha khona ukuthetha "
          "ukuthi inkosi uThixo imfundo umsebenzi ilizwe isizwe abafundisi ukubhala incwadi "
          "iphephandaba urhulumente ekhaya abantwana umfazi indoda amadoda ukuphila ubomi uthando "
          "amanzi imvula iinkomo umhlaba ixesha usuku unyaka iminyaka inyanga ngokuba kuba ngenxa "
          "ukufunda isikolo izikolo ukuhamba ukubona ukuva ukwazi ulwazi ingxoxo umbuzo impendulo "
          "ibandla icawa ukuthandaza umthandazo ukukhanya uxolo umthetho imithetho inkundla ityala "
          "imali ukuthenga ukuthengisa ivenkile ukubhabha khumbula thatha phendula").split()
ENGLISH = ["The annual meeting of the association was held in the town hall on Monday last.",
           "Subscribers are reminded that payment for the current year is now due.",
           "The editor thanks the many correspondents who have written from the districts.",
           "A public notice regarding the registration of voters appears below.",
           "Prices of wool and grain remained steady at the market during the week."]
ADVERT = ("VENKILE ENTSHA ivenkile ithengisa iingubo imali encinane ukuthenga kwaye ukuthengisa "
          "iinkomo amanzi isikolo incwadi abantu bonke bayamenywa ngoku kakhulu")


def historical(word: str, rng: random.Random, p: float = 0.7) -> str:
    """Apply older spelling (inverse ASP, VEL, BH) with probability p."""
    if rng.random() > p:
        return word
    w = word.replace("rh", "r").replace("bh", "b")
    for c in "ptk":
        w = w.replace(c + "h", c)
    return w


def sentence(rng: random.Random, n: int, hist: bool) -> str:
    ws = [rng.choice(MODERN) for _ in range(n)]
    if hist:
        ws = [historical(w, rng) for w in ws]
    s = " ".join(ws)
    return s[0].upper() + s[1:] + "."


def make_edition(path: Path, masthead: str, rng: random.Random) -> None:
    doc = Document()
    doc.add_paragraph(masthead)
    doc.add_paragraph("")
    for _ in range(14):
        doc.add_paragraph(" ".join(sentence(rng, rng.randint(9, 14), True) for _ in range(4)))
    doc.add_paragraph(" ".join(rng.sample(ENGLISH, 3)) + " " + " ".join(rng.sample(ENGLISH, 3)))
    doc.add_paragraph(ADVERT)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = " ".join(sentence(rng, 12, True) for _ in range(3))
    table.cell(0, 1).text = " ".join(sentence(rng, 12, True) for _ in range(3))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> None:
    args = base_parser(__doc__).parse_args()
    cfg = load_config(args.config, smoke=True)
    rng = random.Random(13)
    editions = {"isigidimi": [("Isigidimi_1880_01_01.docx", "ISIGIDIMI SAMA-XOSA. LOVEDALE, JANUARY 1, 1880."),
                              ("Isigidimi_1880_02_02.docx", "ISIGIDIMI SAMA-XOSA. LOVEDALE, FEBRUARY 2, 1880."),
                              ("Isigidimi_1880_03_01.docx", "ISIGIDIMI SAMA-XOSA. LOVEDALE, MARCH 1, 1880."),
                              ("Isigidimi_1880_04_01.docx", "ISIGIDIMI SAMA-XOSA. LOVEDALE, APRIL 1, 1880."),
                              ("Isigidimi_1880_05_01.docx", "ISIGIDIMI SAMA-XOSA. LOVEDALE, MAY 1, 1880.")],
                "imvo": [("Imvo_1885_01_07.docx", "IMVO ZABANTSUNDU. KING WILLIAM'S TOWN, JANUARY 7, 1885."),
                         ("Imvo_1885_01_14.docx", "IMVO ZABANTSUNDU. KING WILLIAM'S TOWN, JANUARY 14, 1885."),
                         ("Imvo_1885_01_21.docx", "IMVO ZABANTSUNDU. KING WILLIAM'S TOWN, JANUARY 21, 1885."),
                         ("Imvo_1885_01_28.docx", "IMVO ZABANTSUNDU. KING WILLIAM'S TOWN, JANUARY 28, 1885."),
                         ("Imvo_1885_02_04.docx", "IMVO ZABANTSUNDU. KING WILLIAM'S TOWN, FEBRUARY 4, 1885.")]}
    for source, eds in editions.items():
        for fname, mast in eds:
            make_edition(P(cfg, "raw", cfg["historical"]["sources"][source], fname), mast, rng)

    rows = []
    for source, splits in (("M1", {"train": 20, "eval": 10, "test": 10}),
                           ("M2", {"validation": 20, "test": 10})):
        for split, n in splits.items():
            for i in range(n):
                text = " ".join(sentence(rng, rng.randint(10, 14), False) for _ in range(12))
                rows.append({"source": source, "doc_id": f"{source}-{split}-{i}", "split": split, "text": text})
    lid = [{"text": " ".join(sentence(rng, 12, False) for _ in range(8)), "label": "xho", "split": "train"}
           for _ in range(60)]
    lid += [{"text": " ".join(rng.sample(ENGLISH, 5)), "label": "eng", "split": "train"} for _ in range(60)]
    P(cfg, "interim").mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_parquet(P(cfg, "interim", "modern_raw.parquet"), index=False)
    pd.DataFrame(lid).to_parquet(P(cfg, "interim", "vuk_langid.parquet"), index=False)
    print("synthetic smoke data written to", P(cfg, "raw").parent, "-- now run: bash run_all.sh --smoke")


if __name__ == "__main__":
    main()
